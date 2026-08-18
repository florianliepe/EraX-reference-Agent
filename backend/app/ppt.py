from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation

from .models import ReferenceDraft
from .structuring import LIMITS, compact


TEMPLATE_PATH = Path(__file__).parent / "assets" / "reference-template.pptx"
TITLE_SHAPE = "Title 12"
METADATA_TABLE = "Table 5"
CONTENT_TABLE = "Table 13"
PICTURE_SHAPE = "Bildplatzhalter 5"
SLIDE_NUMBER_SHAPE = "Slide Number Placeholder 2"


def _shape(slide, name: str):
    try:
        return next(item for item in slide.shapes if item.name == name)
    except StopIteration as exc:
        raise ValueError(f"Template element is missing: {name}") from exc


def _set_paragraph_text(paragraph, value: str) -> None:
    """Replace copy while retaining the inherited paragraph/run formatting."""
    if paragraph.runs:
        first_run = paragraph.runs[0]
        first_run.text = value
        for child in list(paragraph._p):
            if child is paragraph._p.get_or_add_pPr() or child is first_run._r:
                continue
            paragraph._p.remove(child)
    else:
        paragraph.text = value


def _trim_paragraphs(text_frame, count: int) -> None:
    for paragraph in list(text_frame.paragraphs[count:]):
        text_frame._txBody.remove(paragraph._p)


def _set_text_shape(shape, value: str) -> None:
    frame = shape.text_frame
    _set_paragraph_text(frame.paragraphs[0], value)
    _trim_paragraphs(frame, 1)


def _set_labeled_cell(cell, label: str, value: str) -> None:
    paragraph = cell.text_frame.paragraphs[0]
    runs = paragraph.runs
    if len(runs) >= 2:
        runs[0].text = label
        runs[1].text = f": {value}"
        for run in list(runs[2:]):
            paragraph._p.remove(run._r)
    else:
        paragraph.text = f"{label}: {value}"
        paragraph.runs[0].font.bold = True
    _trim_paragraphs(cell.text_frame, 1)


def _sentences(value: str, limit: int) -> list[str]:
    parts = [part.strip() for part in re.split(r"(?<=[.!?])\s+", value) if part.strip()]
    return parts[:limit] or ["Insufficient source evidence"]


def _set_section_cell(cell, heading: str, value: str, kpi_summary: str | None = None) -> None:
    frame = cell.text_frame
    paragraphs = list(frame.paragraphs)
    _set_paragraph_text(paragraphs[0], heading)
    slots = max(1, len(paragraphs) - 1)
    body = _sentences(value, slots)
    if kpi_summary:
        metric_line = compact(f"Key KPIs: {kpi_summary}", 180)
        if slots == 1:
            body = [compact(f"{body[0]} {metric_line}", 300)]
        else:
            body = body[: slots - 1] + [metric_line]
    for index, sentence in enumerate(body, 1):
        target = paragraphs[min(index, len(paragraphs) - 1)]
        _set_paragraph_text(target, sentence)
    _trim_paragraphs(frame, len(body) + 1)


def _replace_picture(slide, image_path: Path | None) -> None:
    if not image_path:
        return
    picture = _shape(slide, PICTURE_SHAPE)
    left, top, width, height = picture.left, picture.top, picture.width, picture.height
    picture._element.getparent().remove(picture._element)
    replacement = slide.shapes.add_picture(str(image_path), left, top, width, height)
    replacement.name = PICTURE_SHAPE
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    frame_ratio = width / height
    if image_ratio > frame_ratio:
        crop = (1 - frame_ratio / image_ratio) / 2
        replacement.crop_left = crop
        replacement.crop_right = crop
    else:
        crop = (1 - image_ratio / frame_ratio) / 2
        replacement.crop_top = crop
        replacement.crop_bottom = crop


def render_ppt(
    draft: ReferenceDraft,
    classification: str,
    output: Path,
    job_id: str,
    model_version: str,
    image_path: Path | None = None,
) -> None:
    if not TEMPLATE_PATH.exists():
        raise ValueError("Approved Eraneos reference template is not installed")

    deck = Presentation(TEMPLATE_PATH)
    if len(deck.slides) != 1:
        raise ValueError("Template invariant violated: expected exactly one source slide")
    slide = deck.slides[0]

    _set_text_shape(_shape(slide, TITLE_SHAPE), compact(draft.title.value, LIMITS["title"]))
    _set_text_shape(_shape(slide, SLIDE_NUMBER_SHAPE), "1")

    metadata = _shape(slide, METADATA_TABLE).table
    _set_labeled_cell(metadata.cell(0, 0), "Client", draft.client.value)
    _set_labeled_cell(metadata.cell(0, 1), "Year", draft.date.value)
    _set_labeled_cell(metadata.cell(1, 0), "Industry", draft.industry.value)
    _set_labeled_cell(metadata.cell(1, 1), "Service", draft.service.value)

    content = _shape(slide, CONTENT_TABLE).table
    _set_section_cell(content.cell(0, 0), "Initial Situation", draft.situation_challenge.value)
    _set_section_cell(content.cell(0, 1), "Approach and Solution", draft.approach.value)
    kpi_summary = "; ".join(f"{item.name}: {item.value}" for item in draft.kpis[:3])
    _set_section_cell(content.cell(1, 0), "Impact", draft.outcome_impact.value, kpi_summary or None)
    _replace_picture(slide, image_path)

    deck.core_properties.title = draft.title.value
    deck.core_properties.subject = "Eraneos client reference one-pager"
    deck.core_properties.comments = (
        f"classification={classification}; job={job_id}; generator={model_version}; "
        "template=Reference Template.pptx"
    )
    if len(deck.slides) != 1:
        raise ValueError("Renderer invariant violated: output must contain exactly one slide")
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
