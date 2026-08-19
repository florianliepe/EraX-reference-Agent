from __future__ import annotations

import re
import time
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import fitz
import httpx
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from .config import settings


SUPPORTED = {".ppt", ".pptx", ".doc", ".docx", ".xls", ".xlsx", ".pdf", ".png", ".jpg", ".jpeg"}
OOXML_SIGNATURE = b"PK"
PDF_SIGNATURE = b"%PDF"
IMAGE_SIGNATURES = (b"\x89PNG", b"\xff\xd8\xff")
CLASSIFICATION_ORDER = {"public": 0, "internal": 1, "confidential": 2, "strict": 3}


@dataclass
class Chunk:
    text: str
    location: str | None = None
    block_type: str = "paragraph"
    page: int | None = None
    bbox: tuple[float, float, float, float] | None = None
    heading_path: tuple[str, ...] = ()
    font_size: float | None = None
    is_bold: bool = False
    repeated: bool = False
    parser: str = "local"
    table_cells: tuple[str, ...] = ()


def validate_signature(path: Path) -> None:
    head = path.read_bytes()[:8]
    suffix = path.suffix.lower()
    if suffix in {".pptx", ".docx", ".xlsx"} and not head.startswith(OOXML_SIGNATURE):
        raise ValueError("File content does not match its OOXML extension")
    if suffix == ".pdf" and not head.startswith(PDF_SIGNATURE):
        raise ValueError("File content does not match PDF extension")
    if suffix in {".png", ".jpg", ".jpeg"} and not any(head.startswith(sig) for sig in IMAGE_SIGNATURES):
        raise ValueError("File content does not match image extension")


def parse_docx(path: Path) -> list[Chunk]:
    document = Document(path)
    return [Chunk(p.text.strip(), f"paragraph {i + 1}") for i, p in enumerate(document.paragraphs) if p.text.strip()]


def parse_pptx(path: Path) -> list[Chunk]:
    deck = Presentation(path)
    chunks: list[Chunk] = []
    for index, slide in enumerate(deck.slides, 1):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                chunks.append(Chunk(shape.text.strip(), f"slide {index}"))
    return chunks


def parse_xlsx(path: Path) -> list[Chunk]:
    book = load_workbook(path, read_only=True, data_only=True)
    chunks: list[Chunk] = []
    for sheet in book.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            text = " | ".join(str(value).strip() for value in row if value not in (None, ""))
            if text:
                chunks.append(Chunk(text, f"{sheet.title}!{row_index}"))
    return chunks


def _normalized_signature(text: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text).strip().casefold())[:240]


def _block_text(block: dict) -> tuple[str, float | None, bool]:
    lines: list[str] = []
    sizes: list[float] = []
    bold = False
    for line in block.get("lines", []):
        parts = []
        for span in line.get("spans", []):
            value = str(span.get("text", ""))
            if value:
                parts.append(value)
            if span.get("size") is not None:
                sizes.append(float(span["size"]))
            bold = bold or "bold" in str(span.get("font", "")).casefold()
        joined = "".join(parts).strip()
        if joined:
            lines.append(joined)
    text = "\n".join(lines)
    text = re.sub(r"(?<=\w)-\n(?=\w)", "", text)
    return text.strip(), max(sizes) if sizes else None, bold


def _looks_like_heading(text: str, font_size: float | None, bold: bool, median_size: float) -> bool:
    compact = re.sub(r"\s+", " ", text).strip()
    return bool(
        compact
        and len(compact) <= 140
        and (
            (bold and (font_size or 0) >= median_size)
            or (font_size or 0) >= median_size * 1.18
            or re.match(r"^(?:\d+[.)]|[A-ZÄÖÜ][^.!?]{2,80}:)$", compact)
        )
    )


def _parse_pdf_local(path: Path) -> list[Chunk]:
    document = fitz.open(path)
    raw: list[dict] = []
    font_sizes: list[float] = []
    table_regions: dict[int, list[fitz.Rect]] = {}
    table_rows: list[Chunk] = []
    for page_index, page in enumerate(document, 1):
        regions: list[fitz.Rect] = []
        try:
            for table_index, table in enumerate(page.find_tables().tables, 1):
                regions.append(fitz.Rect(table.bbox))
                for row_index, row in enumerate(table.extract(), 1):
                    cells = tuple(re.sub(r"\s+", " ", str(value or "")).strip() for value in row)
                    if any(cells):
                        table_rows.append(Chunk(
                            text=" | ".join(value for value in cells if value),
                            location=f"page {page_index}, table {table_index}, row {row_index}",
                            block_type="table_row",
                            page=page_index,
                            bbox=tuple(table.bbox),
                            parser="pymupdf-table",
                            table_cells=cells,
                        ))
        except Exception:
            regions = []
        table_regions[page_index] = regions
        page_height = max(float(page.rect.height), 1)
        for order, block in enumerate(page.get_text("dict", sort=True).get("blocks", [])):
            if block.get("type") != 0:
                continue
            text, font_size, bold = _block_text(block)
            if not text:
                continue
            bbox = tuple(float(value) for value in block.get("bbox", (0, 0, 0, 0)))
            rect = fitz.Rect(bbox)
            if any((rect & region).get_area() >= rect.get_area() * 0.55 for region in regions if rect.get_area()):
                continue
            if font_size:
                font_sizes.append(font_size)
            raw.append({
                "text": text, "page": page_index, "bbox": bbox, "font_size": font_size,
                "bold": bold, "order": order, "height": page_height,
            })
    median_size = sorted(font_sizes)[len(font_sizes) // 2] if font_sizes else 10.0
    margin_signatures = Counter(
        _normalized_signature(item["text"])
        for item in raw
        if item["bbox"][1] <= item["height"] * 0.12 or item["bbox"][3] >= item["height"] * 0.86
    )
    repeat_threshold = max(2, (len(document) + 1) // 2)
    chunks: list[Chunk] = []
    current_heading: str | None = None
    for item in sorted(raw, key=lambda value: (value["page"], value["order"])):
        signature = _normalized_signature(item["text"])
        at_top = item["bbox"][1] <= item["height"] * 0.12
        at_bottom = item["bbox"][3] >= item["height"] * 0.86
        repeated = (at_top or at_bottom) and margin_signatures[signature] >= repeat_threshold
        if repeated:
            block_type = "header" if at_top else "footer"
        elif _looks_like_heading(item["text"], item["font_size"], item["bold"], median_size):
            block_type = "title" if item["page"] == 1 and (item["font_size"] or 0) > median_size * 1.35 else "heading"
            current_heading = re.sub(r"\s+", " ", item["text"]).strip()
        elif re.match(r"^\s*(?:[•▪◦-]|\d+[.)]|[ivx]+[.)])\s+", item["text"], re.I):
            block_type = "list_item"
        elif re.search(r"\b(?:mit freundlichen grüßen|geschäftsführer)\b", item["text"], re.I):
            block_type = "signature"
        else:
            block_type = "paragraph"
        chunks.append(Chunk(
            text=item["text"],
            location=f"page {item['page']}",
            block_type=block_type,
            page=item["page"],
            bbox=item["bbox"],
            heading_path=(current_heading,) if current_heading and block_type not in {"heading", "title"} else (),
            font_size=item["font_size"],
            is_bold=item["bold"],
            repeated=repeated,
            parser="pymupdf-layout",
        ))
    chunks.extend(table_rows)
    document.close()
    return sorted(chunks, key=lambda item: (item.page or 0, item.bbox[1] if item.bbox else 0, item.location or ""))


def needs_document_intelligence(chunks: list[Chunk], page_count: int) -> bool:
    usable = [item for item in chunks if not item.repeated and item.block_type not in {"header", "footer"}]
    return sum(len(item.text) for item in usable) < max(120, page_count * 80)


def _parse_pdf_azure(path: Path) -> list[Chunk]:
    endpoint = (settings.azure_document_intelligence_endpoint or "").rstrip("/")
    if not endpoint or not settings.azure_document_intelligence_api_key:
        return []
    url = (
        f"{endpoint}/documentintelligence/documentModels/prebuilt-layout:analyze"
        f"?api-version={settings.azure_document_intelligence_api_version}"
    )
    headers = {
        "Ocp-Apim-Subscription-Key": settings.azure_document_intelligence_api_key,
        "Content-Type": "application/octet-stream",
    }
    response = httpx.post(url, headers=headers, content=path.read_bytes(), timeout=30)
    response.raise_for_status()
    operation = response.headers.get("operation-location")
    if not operation:
        raise ValueError("Azure Document Intelligence did not return an operation URL")
    result = None
    for _ in range(30):
        poll = httpx.get(operation, headers={"Ocp-Apim-Subscription-Key": settings.azure_document_intelligence_api_key}, timeout=15)
        poll.raise_for_status()
        payload = poll.json()
        if payload.get("status") == "succeeded":
            result = payload.get("analyzeResult", {})
            break
        if payload.get("status") == "failed":
            raise ValueError("Azure Document Intelligence layout analysis failed")
        time.sleep(1)
    if result is None:
        raise TimeoutError("Azure Document Intelligence layout analysis timed out")
    chunks: list[Chunk] = []
    for order, paragraph in enumerate(result.get("paragraphs", []), 1):
        regions = paragraph.get("boundingRegions") or []
        page = regions[0].get("pageNumber") if regions else None
        role = paragraph.get("role")
        block_type = "heading" if role in {"title", "sectionHeading"} else "paragraph"
        chunks.append(Chunk(
            text=str(paragraph.get("content", "")).strip(),
            location=f"page {page}" if page else f"paragraph {order}",
            page=page,
            block_type=block_type,
            parser="azure-document-intelligence",
        ))
    for table_index, table in enumerate(result.get("tables", []), 1):
        by_row: dict[int, list[tuple[int, str]]] = {}
        for cell in table.get("cells", []):
            by_row.setdefault(int(cell.get("rowIndex", 0)), []).append((int(cell.get("columnIndex", 0)), str(cell.get("content", ""))))
        page = (table.get("boundingRegions") or [{}])[0].get("pageNumber")
        for row_index, cells in sorted(by_row.items()):
            values = tuple(value for _, value in sorted(cells))
            chunks.append(Chunk(
                text=" | ".join(values), location=f"page {page}, table {table_index}, row {row_index + 1}",
                page=page, block_type="table_row", parser="azure-document-intelligence", table_cells=values,
            ))
    return [item for item in chunks if item.text]


def parse_pdf(path: Path) -> list[Chunk]:
    local = _parse_pdf_local(path)
    with fitz.open(path) as document:
        page_count = len(document)
    if needs_document_intelligence(local, page_count):
        azure = _parse_pdf_azure(path)
        if azure:
            return azure
    return local


def parse_image(path: Path) -> list[Chunk]:
    try:
        import pytesseract
        text = pytesseract.image_to_string(Image.open(path)).strip()
    except Exception as exc:
        raise ValueError("OCR unavailable or image contains no readable text") from exc
    return [Chunk(text, "OCR")] if text else []


PARSERS: dict[str, Callable[[Path], list[Chunk]]] = {
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".xlsx": parse_xlsx,
    ".pdf": parse_pdf,
    ".png": parse_image,
    ".jpg": parse_image,
    ".jpeg": parse_image,
}


def parse_file(path: Path) -> list[Chunk]:
    if path.suffix.lower() in {".doc", ".xls", ".ppt"}:
        raise ValueError("Legacy Office format requires LibreOffice conversion to OOXML")
    validate_signature(path)
    chunks = PARSERS[path.suffix.lower()](path)
    normalized: list[Chunk] = []
    for chunk in chunks:
        if not chunk.text.strip():
            continue
        chunk.text = re.sub(r"[ \t]+", " ", chunk.text).strip()
        normalized.append(chunk)
    return normalized


def detect_classification(chunks: list[Chunk]) -> str:
    detected = "public"
    patterns = {
        "strict": r"\b(strictly confidential|strict)\b",
        "confidential": r"\bconfidential\b",
        "internal": r"\b(internal use only|internal)\b",
        "public": r"\bpublic\b",
    }
    joined = " ".join(chunk.text for chunk in chunks).lower()
    for label, pattern in patterns.items():
        if re.search(pattern, joined) and CLASSIFICATION_ORDER[label] > CLASSIFICATION_ORDER[detected]:
            detected = label
    return detected
