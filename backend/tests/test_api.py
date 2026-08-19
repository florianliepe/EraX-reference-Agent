from io import BytesIO
import json
from docx import Document
from fastapi.testclient import TestClient
from backend.app.config import settings
from backend.app.main import FAILED_AUTH_LIMIT, app, attempts


def test_full_flow(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    document = Document()
    document.add_paragraph("Client: Contoso Energy")
    document.add_paragraph("Project title: Grid operations redesign")
    document.add_paragraph("Challenge: Fragmented processes delayed incident response.")
    document.add_paragraph("Approach: Redesigned workflows and introduced a governed analytics platform.")
    document.add_paragraph("Outcome: Incident triage time reduced by 25 percent.")
    payload = BytesIO(); document.save(payload)
    client = TestClient(app); headers = {"X-Pilot-Password": settings.pilot_password}
    assert client.get("/auth/check", headers=headers).json() == {"authenticated": True}
    assert client.get("/auth/check", headers={"X-Pilot-Password": "wrong"}).status_code == 401
    uploaded = client.post("/upload", files={"files": ("reference.docx", payload.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=headers)
    assert uploaded.status_code == 200
    started = client.post(f"/generate/{uploaded.json()['session_id']}", headers=headers)
    job_id = started.json()["job_id"]
    result = client.get(f"/result/{job_id}", headers=headers)
    assert result.status_code == 200
    assert result.json()["fields"]["client"]["value"] == "Contoso Energy"
    assert result.json()["fields"]["kpis"][0]["value"] == "25 percent"
    kpi_id = result.json()["fields"]["kpis"][0]["id"]
    patched = client.patch(
        f"/result/{job_id}",
        headers=headers,
        json={"values": {}, "kpis": [
            {"id": kpi_id, "name": "Incident triage reduction", "value": "25%"},
            {"name": "Reviewer-supplied target", "value": "42%"},
        ]},
    )
    assert patched.json()["fields"]["kpis"][0]["edited"] is True
    assert patched.json()["fields"]["kpis"][0]["evidence"][0]["source_file"] == "reference.docx"
    assert client.post(
        f"/approve/{job_id}",
        headers=headers,
        json={"approved_by": "", "reuse_allowed": False},
    ).status_code == 422
    approval = client.post(
        f"/approve/{job_id}",
        headers=headers,
        json={"approved_by": "Pilot Reviewer", "reuse_allowed": True},
    )
    assert approval.status_code == 200
    assert approval.json()["publication"]["status"] == "local"
    knowledge = client.get(f"/knowledge/{job_id}", headers=headers)
    assert knowledge.status_code == 200
    graph = json.loads(knowledge.content)["@graph"]
    approved_claims = [item for item in graph if item.get("@type") == "erax:Claim"]
    assert any(item["schema:description"] == "Incident triage reduction: 25%" for item in approved_claims)
    assert all(item["erax:status"] == "human-approved" for item in approved_claims)
    ungrounded = next(item for item in approved_claims if item["schema:description"] == "Reviewer-supplied target: 42%")
    assert ungrounded["prov:wasDerivedFrom"] == []
    assert ungrounded["erax:reuseAllowed"] is False
    download = client.get(f"/download/{job_id}", headers=headers)
    assert download.status_code == 200


def test_valid_status_polling_is_not_rate_limited():
    attempts.clear()
    client = TestClient(app)
    headers = {"X-Pilot-Password": settings.pilot_password}
    responses = [client.get("/auth/check", headers=headers) for _ in range(75)]
    assert all(response.status_code == 200 for response in responses)
    assert not attempts


def test_only_failed_authentication_uses_login_rate_limit():
    attempts.clear()
    client = TestClient(app)
    headers = {"X-Pilot-Password": "wrong"}
    for _ in range(FAILED_AUTH_LIMIT):
        assert client.get("/auth/check", headers=headers).status_code == 401
    limited = client.get("/auth/check", headers=headers)
    assert limited.status_code == 429
    assert int(limited.headers["Retry-After"]) >= 1
    attempts.clear()
